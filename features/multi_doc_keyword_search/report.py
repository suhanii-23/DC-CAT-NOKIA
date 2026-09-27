"""The Multi-Document Keyword Search Report: one sheet, one row per match.

Why this does not go through common/excel.py
--------------------------------------------
common/excel.py is the shared report for the per-document analysis
features. It always writes a Summary sheet and always starts every
feature sheet with page, severity, message and confidence columns. That
suits findings that are faults, but a keyword search reports locations,
not faults. A "severity" or "confidence" column there is noise to a
documentation user. Changing the shared writer would change every other
feature's report, so this feature writes its own single-sheet workbook
with the same library (openpyxl) instead.

app/cli.py still reports this feature through common/excel.py, alongside
the other features, and that path is unchanged.
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from features.multi_doc_keyword_search.service import MultiDocSearchResult

SHEET_TITLE = "Keyword Search"

# (header, column width in characters). The order is the column order.
COLUMNS = [
    ("Document", 32),
    ("Page", 8),
    ("Keyword", 18),
    ("Match", 18),
    ("Context", 100),
]

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="124191")  # dark blue
_TOP = Alignment(vertical="top")
_WRAPPED = Alignment(vertical="top", wrap_text=True)


def write_xlsx(result: MultiDocSearchResult, out_path: str) -> None:
    """Write every match in `result` to `out_path` as one sheet.

    A search with no matches still produces a valid workbook: the header
    row and nothing under it.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE

    sheet.append([header for header, _width in COLUMNS])
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _TOP

    for row in result.rows():
        sheet.append([row.document, row.page, row.keyword, row.match, row.context])

    for index, (_header, width) in enumerate(COLUMNS, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = width
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = _TOP
        cells[-1].alignment = _WRAPPED  # Context is the long one

    sheet.freeze_panes = "A2"  # header stays visible while scrolling
    sheet.auto_filter.ref = sheet.dimensions  # filter by document or page

    folder = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(folder, exist_ok=True)
    workbook.save(out_path)
