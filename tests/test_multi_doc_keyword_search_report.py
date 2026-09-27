"""End-to-end tests: N documents + 1 keyword -> 1 clean XLSX report.

Real PDF files are written to disk and searched through the path-based
entry point and the feature's own CLI, exactly as a user runs it:

    python -m features.multi_doc_keyword_search.cli \
        document1.pdf document2.pdf document3.pdf --keyword embedded --excel out.xlsx

The corpus (built per test):

    document1.pdf  3 pages  p2: "embedded" x2, p3: "EMBEDDED" x1, "firmware" on p2
    document2.pdf  2 pages  p2: "embedded" x1
    document3.pdf  1 page   no keyword at all
"""
from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

fitz = pytest.importorskip("pymupdf")

from features.multi_doc_keyword_search import cli  # noqa: E402
from features.multi_doc_keyword_search.matcher import (  # noqa: E402
    compile_keyword,
    find_occurrences,
)
from features.multi_doc_keyword_search.report import COLUMNS  # noqa: E402
from features.multi_doc_keyword_search.service import (  # noqa: E402
    MultiDocKeywordSearchService,
)

EXPECTED_HEADER = ("Document", "Page", "Keyword", "Match", "Context")

UNWANTED_COLUMNS = {
    "severity", "confidence", "priority", "finding type", "bug type",
    "status", "suggested fix", "remediation", "similarity score",
    "semantic score", "agent", "reason", "risk", "message",
}


def _pdf(path: Path, pages: list[str]) -> str:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    doc.save(str(path))
    doc.close()
    return str(path)


@pytest.fixture
def corpus(tmp_path) -> list[str]:
    return [
        _pdf(
            tmp_path / "document1.pdf",
            [
                "Introduction to the course.",
                "The embedded controller starts first.\n"
                "Embedded firmware is updated later.",
                "EMBEDDED systems are small.",
            ],
        ),
        _pdf(
            tmp_path / "document2.pdf",
            ["Nothing relevant here.", "An embedded board runs the loop."],
        ),
        _pdf(tmp_path / "document3.pdf", ["Topology diagram only."]),
    ]


def _search(paths, keyword="embedded"):
    return MultiDocKeywordSearchService().search(paths, keyword)


def _report(tmp_path, argv) -> tuple[int, list[tuple]]:
    out = tmp_path / "out.xlsx"
    code = cli.main([*argv, "--excel", str(out)])
    rows = list(load_workbook(out)["Keyword Search"].iter_rows(values_only=True))
    return code, rows


# --- Test 1: three documents, all searched ------------------------------


def test_all_three_documents_are_searched(corpus):
    result = _search(corpus)

    assert result.status == "ok"
    assert result.documents_searched == 3
    assert result.errors == []
    assert sorted(result.document_names().values()) == [
        "document1.pdf", "document2.pdf", "document3.pdf",
    ]


# --- Test 2: keyword in one document only -------------------------------


def test_a_keyword_in_one_document_only_reports_only_that_document(corpus):
    rows = _search(corpus, "firmware").rows()

    assert [row.document for row in rows] == ["document1.pdf"]
    assert rows[0].page == 2


# --- Test 3: keyword in several documents -------------------------------


def test_results_from_every_matching_document_are_included(corpus):
    result = _search(corpus)

    assert {row.document for row in result.rows()} == {"document1.pdf", "document2.pdf"}
    assert result.documents_with_matches == 2


# --- Test 4: several occurrences -> one row each -------------------------


def test_every_occurrence_is_its_own_row(corpus):
    rows = _search(corpus).rows()

    per_document = {}
    for row in rows:
        per_document[row.document] = per_document.get(row.document, 0) + 1
    assert per_document == {"document1.pdf": 3, "document2.pdf": 1}
    assert len(rows) == 4


def test_two_occurrences_in_one_sentence_are_two_rows():
    pattern = compile_keyword("embedded")
    found = find_occurrences("Embedded code on embedded boards.", pattern)
    assert [o.text for o in found] == ["Embedded", "embedded"]


# --- Test 5: case-insensitive --------------------------------------------


def test_every_casing_in_the_document_is_found_and_shown_as_written(corpus):
    rows = _search(corpus).rows()
    assert [row.match for row in rows] == ["embedded", "Embedded", "EMBEDDED", "embedded"]


@pytest.mark.parametrize("keyword", ["embedded", "Embedded", "EMBEDDED"])
def test_the_casing_of_the_keyword_does_not_change_the_result(corpus, keyword):
    result = _search(corpus, keyword)
    assert result.total_matches == 4
    assert {row.keyword for row in result.rows()} == {keyword}


def test_whole_word_contract_is_kept(corpus, tmp_path):
    # The established contract (tests/test_multi_doc_keyword_search.py):
    # whole words only, so a longer word is not a match.
    path = _pdf(tmp_path / "words.pdf", ["embeddedness and embeddedsystem, embedded."])
    assert [row.match for row in _search([path]).rows()] == ["embedded"]


# --- Test 6: page numbers ------------------------------------------------


def test_each_row_carries_the_page_it_was_found_on(corpus):
    rows = _search(corpus).rows()
    assert [(row.document, row.page) for row in rows] == [
        ("document1.pdf", 2),
        ("document1.pdf", 2),
        ("document1.pdf", 3),
        ("document2.pdf", 2),
    ]


# --- Test 7: document attribution ----------------------------------------


def test_rows_name_the_source_file_not_a_path(corpus):
    for row in _search(corpus).rows():
        assert row.document in {"document1.pdf", "document2.pdf"}
        assert os.sep not in row.document


def test_two_files_with_the_same_name_keep_their_full_paths(tmp_path):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    first = _pdf(tmp_path / "a" / "manual.pdf", ["embedded one"])
    second = _pdf(tmp_path / "b" / "manual.pdf", ["embedded two"])

    documents = [row.document for row in _search([first, second]).rows()]

    assert documents == [os.path.normpath(first), os.path.normpath(second)]


# --- Test 8: context -----------------------------------------------------


def test_context_is_the_surrounding_sentence(corpus):
    rows = _search(corpus).rows()
    assert rows[0].context == "The embedded controller starts first."
    assert rows[1].context == "Embedded firmware is updated later."
    assert rows[3].context == "An embedded board runs the loop."


def test_a_long_run_of_text_gets_a_trimmed_window_not_a_whole_page():
    filler = "alpha beta gamma delta " * 40
    pattern = compile_keyword("embedded")
    context = find_occurrences(filler + "embedded " + filler, pattern)[0].context

    assert context.startswith("...") and context.endswith("...")
    assert "embedded" in context
    assert len(context) < 250
    # Trimmed to whole words: no fragment of a word at either cut.
    assert set(context.strip(".").split()) <= {"alpha", "beta", "gamma", "delta", "embedded"}


# --- Test 9: no matches --------------------------------------------------


def test_no_matches_anywhere_still_writes_a_valid_report(corpus, tmp_path, capsys):
    code, rows = _report(tmp_path, [*corpus, "--keyword", "authentication"])

    assert code == 0
    assert rows == [EXPECTED_HEADER]  # headers, zero result rows
    assert "No matches found for keyword: authentication" in capsys.readouterr().out


# --- Test 10: the XLSX ---------------------------------------------------


def test_report_has_exactly_the_five_columns_and_every_match(corpus, tmp_path):
    code, rows = _report(tmp_path, [*corpus, "--keyword", "embedded"])

    assert code == 0
    header, data = rows[0], rows[1:]
    assert header == EXPECTED_HEADER
    assert not {str(h).lower() for h in header} & UNWANTED_COLUMNS
    assert data == [
        ("document1.pdf", 2, "embedded", "embedded", "The embedded controller starts first."),
        ("document1.pdf", 2, "embedded", "Embedded", "Embedded firmware is updated later."),
        ("document1.pdf", 3, "embedded", "EMBEDDED", "EMBEDDED systems are small."),
        ("document2.pdf", 2, "embedded", "embedded", "An embedded board runs the loop."),
    ]


def test_report_is_one_readable_sheet(corpus, tmp_path):
    out = tmp_path / "out.xlsx"
    cli.main([*corpus, "--keyword", "embedded", "--excel", str(out)])
    workbook = load_workbook(out)

    assert workbook.sheetnames == ["Keyword Search"]  # no summary or extra sheets
    sheet = workbook["Keyword Search"]
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].font.bold
    assert sheet["E2"].alignment.wrap_text  # Context wraps
    widths = [sheet.column_dimensions[c].width for c in "ABCDE"]
    assert widths == [width for _header, width in COLUMNS]


def test_the_report_module_does_not_use_the_shared_findings_columns():
    # common/excel.py always writes severity and confidence; this report
    # must not go through it.
    source = Path(cli.__file__).with_name("report.py").read_text(encoding="utf-8")
    assert "from common import excel" not in source
    assert "write_report" not in source


# --- Test 11: a single document ------------------------------------------


def test_a_single_document_works_through_the_same_interface(corpus, tmp_path):
    code, rows = _report(tmp_path, [corpus[1], "--keyword", "embedded"])

    assert code == 0
    assert rows[1:] == [
        ("document2.pdf", 2, "embedded", "embedded", "An embedded board runs the loop."),
    ]


# --- CLI summary ---------------------------------------------------------


def test_summary_lists_every_document_with_its_count(corpus, tmp_path, capsys):
    out = tmp_path / "out.xlsx"
    cli.main([*corpus, "--keyword", "embedded", "--excel", str(out)])
    printed = capsys.readouterr().out

    assert "Multi-document keyword search completed." in printed
    assert "Keyword: embedded" in printed
    assert "Documents searched: 3" in printed
    assert "document1.pdf: 3 matches" in printed
    assert "document2.pdf: 1 match\n" in printed
    assert "document3.pdf: 0 matches" in printed
    assert "Total matches: 4" in printed
    assert f"Output: {out}" in printed


# --- real-document text handling -----------------------------------------


def test_a_word_hyphenated_across_a_line_break_is_found(tmp_path):
    path = _pdf(tmp_path / "hyphen.pdf", ["The system uses em-\nbedded control."])
    rows = _search([path]).rows()

    assert [row.match for row in rows] == ["embedded"]
    assert rows[0].context == "The system uses em-bedded control."


def test_soft_hyphens_are_matched_through_and_removed_from_the_context():
    soft = "\u00ad"
    text = f"develop{soft}\nments are fast. The em{soft}\nbedded board runs."
    found = find_occurrences(text, compile_keyword("embedded"))

    assert [o.text for o in found] == ["embedded"]
    assert found[0].context == "The embedded board runs."
    # A soft hyphen is part of the word: no whole-word hit inside it.
    assert find_occurrences(text, compile_keyword("develop")) == []


def test_a_scanned_pdf_is_flagged_instead_of_reported_as_a_clean_zero(
    corpus, tmp_path, capsys
):
    scanned = tmp_path / "scanned.pdf"
    doc = fitz.open()
    for _ in range(3):
        page = doc.new_page()
        pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
        page.insert_image(fitz.Rect(72, 72, 300, 300), pixmap=pixmap)
    doc.save(str(scanned))
    doc.close()

    result = _search([corpus[0], str(scanned)])
    flags = {os.path.basename(d.path): d.looks_scanned for d in result.documents}
    assert flags == {"document1.pdf": False, "scanned.pdf": True}

    cli.main([corpus[0], str(scanned), "--keyword", "embedded", "--no-excel"])
    assert "scanned.pdf: 0 matches  (almost no extractable text" in capsys.readouterr().out


def test_a_short_text_pdf_is_not_mistaken_for_a_scan(corpus):
    result = _search([corpus[2]])  # one line of text, no image
    assert result.documents[0].looks_scanned is False


# --- performance: nothing heavy is loaded ---------------------------------


def test_the_cli_loads_no_model_or_orchestration_library(corpus, tmp_path):
    heavy = ["torch", "sentence_transformers", "transformers", "faiss", "langgraph"]
    script = (
        "import sys\n"
        "from features.multi_doc_keyword_search import cli\n"
        f"code = cli.main({[*corpus]!r} + ['--keyword', 'embedded', '--no-excel'])\n"
        f"loaded = [m for m in {heavy!r} if m in sys.modules]\n"
        "print('LOADED', loaded)\n"
        "sys.exit(code)\n"
    )
    root = Path(__file__).resolve().parent.parent
    completed = subprocess.run(
        [sys.executable, "-c", script],
        cwd=root,
        capture_output=True,
        text=True,
        timeout=120,
    )

    assert completed.returncode == 0, completed.stderr
    assert "LOADED []" in completed.stdout
