"""Dictionary spell checking, spell_mode selection, and the per-request
terminology Excel (options["terms_excel"])."""
from __future__ import annotations

import pytest

pytest.importorskip("wordfreq")

from openpyxl import Workbook

from common.contracts import Document, Paragraph
from features.spell_check.dictionary import DictionaryChecker
from features.spell_check.service import SpellCheckService
from features.spell_check.tools.terminology import TerminologyTool


@pytest.fixture(scope="module")
def checker() -> DictionaryChecker:
    return DictionaryChecker()


def _doc(*texts: str) -> Document:
    return Document(
        path="x.pdf",
        format="pdf",
        page_count=1,
        paragraphs=[Paragraph(text=t, page=1, index=i) for i, t in enumerate(texts)],
    )


@pytest.mark.parametrize(
    "typo, fix",
    [
        ("recieve", "receive"), ("teh", "the"), ("thier", "their"),
        ("seperate", "separate"), ("definately", "definitely"),
        ("occured", "occurred"), ("paramter", "parameter"),
        ("batery", "battery"), ("lenght", "length"), ("adress", "address"),
    ],
)
def test_dictionary_catches_common_misspellings(checker, typo, fix):
    assert checker.suggest(typo) == fix


@pytest.mark.parametrize(
    "word",
    [
        # Real words that sit one edit from a far more common word -- each
        # was a false positive on a real Nokia/DC CAT guide during tuning.
        "warms", "cabled", "automates", "identifier", "parses", "headers",
        "Standardizes", "checkpointed",
        # Ordinary and short words.
        "handset", "charger", "network", "the", "lte", "ssd",
    ],
)
def test_dictionary_leaves_real_words_alone(checker, word):
    assert checker.suggest(word) is None


def test_dictionary_skips_urls_paths_and_identifiers(checker):
    text = "See www.nokia.com, /etc/confg, v2.1, airscale_rnc and user@nokai.com."
    assert checker.check(text) == []


def test_dictionary_keeps_case(checker):
    [change] = checker.check("Recieve the file.")
    assert (change.original, change.suggestion) == ("Recieve", "Receive")


def test_dictionary_mode_reports_method_and_issue_type():
    result = SpellCheckService().process(
        _doc("The paramter is set."), {"spell_mode": "dictionary"}
    )
    assert result.status == "ok"
    [finding] = result.findings
    assert finding.details["suggested_correction"] == "parameter"
    assert finding.details["issue_type"] == "dictionary_spelling"
    assert finding.details["method"] == "dictionary"
    assert finding.details["corrected_sentence"] == "The parameter is set."
    assert result.meta["spell_mode"] == "dictionary"


def test_unknown_spell_mode_fails_without_raising():
    result = SpellCheckService().process(_doc("Text."), {"spell_mode": "nope"})
    assert result.status == "failed"
    assert "nope" in result.error


def test_terms_excel_suppresses_listed_words(tmp_path):
    doc = _doc("Open the Docset and the batery panel.")
    service = SpellCheckService()
    flagged = {
        f.details["word"]
        for f in service.process(doc, {"spell_mode": "dictionary"}).findings
    }
    assert {"Docset", "batery"} <= flagged

    path = tmp_path / "terms.xlsx"
    wb = Workbook()
    wb.active.append(["Acronym"])
    wb.active.append(["Docset"])
    wb.save(path)

    result = service.process(
        doc, {"spell_mode": "dictionary", "terms_excel": str(path)}
    )
    assert [f.details["word"] for f in result.findings] == ["batery"]
    assert result.meta["uploaded_terms"] == 1


def test_terms_excel_without_header_is_read_as_a_plain_list(tmp_path):
    path = tmp_path / "terms.xlsx"
    wb = Workbook()
    for term in ("Docset", "Network Services Platform"):
        wb.active.append([term])
    wb.save(path)

    tool = TerminologyTool.from_excel(path)
    assert tool.is_protected("Docset")
    assert tool.is_protected("Services")


def test_unreadable_terms_excel_degrades_to_defaults(tmp_path):
    path = tmp_path / "terms.xlsx"
    path.write_text("not an excel file")
    result = SpellCheckService().process(
        _doc("The paramter is set."),
        {"spell_mode": "dictionary", "terms_excel": str(path)},
    )
    assert result.status == "ok"
    assert len(result.findings) == 1
